#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Gera o deck de exemplo "Avaliacao formativa" no design system Cefor (variante Solida),
estrutura Assertion-Evidence (titulo=afirmacao, corpo=evidencia).

Saidas (mesma fonte de dados):
  - avaliacao-formativa.pptx   (python-pptx, editavel)  -> entregavel
  - thumbs/slide-NN.png        (Pillow)                 -> verificacao/figuras

Coordenadas em px sobre um canvas 1280x720; convertidas para polegadas no pptx.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = 1280, 720  # canvas de referencia (16:9)

# --- Paleta Cefor (variante Solida) ---
AZUL   = "2C459A"
OLIVA  = "8C9A0D"
LIMA   = "B0CB1F"
GRENA  = "CC1111"
TEXTO  = "2B2B2B"
TEXTO2 = "5A5A5A"
PAINEL = "EAEAEC"
BRANCO = "FFFFFF"

def rgb(h): return RGBColor.from_string(h)
def pil(h): return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

# --- conversao px -> polegadas (pptx) ---
SCALE_IN = 13.333 / W
def IN(px): return Inches(px * SCALE_IN)

# --- fontes Open Sans / Arial (Windows) ---
def font(size, bold=False):
    paths = [
        r"C:\Windows\Fonts\OpenSans-Bold.ttf" if bold else r"C:\Windows\Fonts\OpenSans-Regular.ttf",
        r"C:\Windows\Fonts\OpenSans.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
    ]
    for p in paths:
        try:
            return ImageFont.truetype(p, size)
        except Exception:
            continue
    return ImageFont.load_default()

# ============================================================
# DADOS DOS SLIDES (assertion-evidence)
# ============================================================
SLIDES = [
    {"kind": "title", "assertion": "Avaliação formativa: o feedback que faz aprender",
     "sub": "Formação docente Cefor   ·   objetivo: produzir feedback que o aluno usa",
     "notes": "Apresenta a sessão e o objetivo: sair daqui produzindo feedback que o aluno usa."},
    {"kind": "timeline", "assertion": "Quando seu feedback chega, ainda dá tempo de mudar a rota?",
     "items": ["Aula", "Entrega", "Nota (tarde)"],
     "notes": "Provoca a experiência do professor: feedback pós-nota raramente muda algo."},
    {"kind": "compare", "assertion": "A formativa acontece durante o percurso; a somativa, no fim",
     "items": [("Formativa", "Ajusta o ensino durante o percurso"),
               ("Somativa", "Certifica o resultado ao final")],
     "notes": "Ambas têm lugar; a formativa é a que muda o ensino em tempo real."},
    {"kind": "compare", "assertion": "Nota informa um resultado; feedback orienta o próximo passo",
     "items": [("Nota: 7,0", "Informa um resultado"),
               ("Feedback", "Orienta o próximo passo")],
     "notes": "O aluno não sabe o que fazer com um número; sabe o que fazer com uma instrução."},
    {"kind": "process", "assertion": "Bom feedback responde a três perguntas",
     "items": [("Onde vou?", "meta — feed up"),
               ("Como estou indo?", "situação — feed back"),
               ("E agora?", "próximo passo — feed forward")],
     "notes": "Feed up, feed back, feed forward (Hattie & Timperley, 2007)."},
    {"kind": "cycle", "assertion": "Feedback fecha a distância entre onde o aluno está e a meta",
     "items": ["Evidência", "Interpretação", "Ação", "Nova evidência"],
     "notes": "A avaliação formativa é um laço de regulação, não um evento único."},
    {"kind": "task", "assertion": "Sua vez: transforme uma nota em feedback acionável",
     "items": ["Pegue um '7,0' e reescreva usando as 3 perguntas. 5 min, em duplas."],
     "notes": "Atividade de 5 min; comparar em duplas."},
    {"kind": "summary", "assertion": "Avaliar para aprender é redirecionar o ensino enquanto há tempo",
     "items": ["Durante o percurso", "Acionável", "Focado e a tempo"],
     "notes": "Retoma o objetivo; foco e tempestividade vencem volume."},
    {"kind": "end", "assertion": "Cefor — Formação docente",
     "sub": "cefor.ifes.edu.br", "notes": ""},
]

# ============================================================
# RENDER 1 — PNG (Pillow), layout fiel
# ============================================================
def wrap(draw, text, fnt, max_w):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= max_w:
            cur = t
        else:
            if cur: lines.append(cur)
            cur = w
    if cur: lines.append(cur)
    return lines

def arrow_png(d, x1, x2, y, color):
    d.line([x1, y, x2 - 10, y], fill=pil(color), width=6)
    d.polygon([(x2 - 12, y - 9), (x2, y), (x2 - 12, y + 9)], fill=pil(color))

def brand_png(d):
    d.rectangle([40, 92, 1240, 648], fill=pil(PAINEL))           # painel de conteudo
    d.rounded_rectangle([40, 40, 250, 150], radius=18, fill=pil(OLIVA))  # bloco oliva topo
    d.line([146, 150, 146, 636], fill=pil(OLIVA), width=3)        # linha-guia
    d.rounded_rectangle([0, 650, 250, 720], radius=18, fill=pil(AZUL))   # bloco azul base
    d.line([700, 690, 1025, 690], fill=pil(OLIVA), width=2)       # rodape (curto, nao cruza o texto)
    d.text((1240, 678), "cefor.ifes.edu.br", font=font(20, True), fill=pil(OLIVA), anchor="ra")

def title_png(d, text, color=AZUL, top=120, left=290, right=1200, size=40):
    fnt = font(size, True)
    y = top
    for ln in wrap(d, text, fnt, right - left):
        d.text((left, y), ln, font=fnt, fill=pil(color)); y += int(size * 1.18)
    return y

def card_png(d, x, y, w, h, head, body, head_color=AZUL, border=None):
    d.rounded_rectangle([x, y, x + w, y + h], radius=14, fill=pil(BRANCO),
                        outline=pil(border) if border else pil("D8D8D8"),
                        width=3 if border else 1)
    d.text((x + 22, y + 18), head, font=font(24, True), fill=pil(head_color))
    if body:
        fnt = font(19)
        yy = y + 58
        for ln in wrap(d, body, fnt, w - 44):
            d.text((x + 22, yy), ln, font=fnt, fill=pil(TEXTO)); yy += 24

def render_png(i, s):
    img = Image.new("RGB", (W, H), pil(BRANCO))
    d = ImageDraw.Draw(img)
    k = s["kind"]

    if k == "title":
        d.rounded_rectangle([40, 40, 250, 150], radius=18, fill=pil(OLIVA))
        d.rounded_rectangle([0, 560, 360, 720], radius=18, fill=pil(AZUL))
        title_png(d, s["assertion"], color=AZUL, top=250, left=120, right=1160, size=52)
        d.text((120, 470), s["sub"], font=font(22), fill=pil(TEXTO2))
        d.text((1240, 678), "cefor.ifes.edu.br", font=font(20, True), fill=pil(OLIVA), anchor="ra")
        img.save(os.path.join(THUMBS, f"slide-{i:02d}.png")); return

    if k == "end":
        d.rectangle([0, 0, W, H], fill=pil(AZUL))
        d.text((W/2, H/2 - 30), s["assertion"], font=font(46, True), fill=pil(BRANCO), anchor="mm")
        d.text((W/2, H/2 + 40), s["sub"], font=font(24), fill=pil(LIMA), anchor="mm")
        img.save(os.path.join(THUMBS, f"slide-{i:02d}.png")); return

    brand_png(d)
    cy = title_png(d, s["assertion"])
    top = max(cy + 24, 270)

    if k == "timeline":
        items = s["items"]; n = len(items); x = 300; w = 250; gap = 70; y = 360
        for j, it in enumerate(items):
            cx = x + j * (w + gap)
            late = (j == n - 1)
            card_png(d, cx, y, w, 110, it, "", head_color=GRENA if late else AZUL,
                     border=GRENA if late else None)
            if j < n - 1:
                arrow_png(d, cx + w + 12, cx + w + gap - 12, y + 55, TEXTO2)

    elif k == "compare":
        (h1, b1), (h2, b2) = s["items"]
        card_png(d, 300, top, 410, 230, h1, b1, head_color=AZUL)
        card_png(d, 770, top, 410, 230, h2, b2, head_color=GRENA)
        d.text((740, top + 115), "x", font=font(34, True), fill=pil(TEXTO2), anchor="mm")

    elif k == "process":
        items = s["items"]; x = 300; w = 270; gap = 35; y = top
        for j, (h1, b1) in enumerate(items):
            cx = x + j * (w + gap)
            d.ellipse([cx, y, cx + 44, y + 44], fill=pil(OLIVA))
            d.text((cx + 22, y + 22), str(j + 1), font=font(24, True), fill=pil(BRANCO), anchor="mm")
            card_png(d, cx, y + 60, w, 150, h1, b1, head_color=AZUL)

    elif k == "cycle":
        items = s["items"]; x = 300; w = 200; gap = 33; y = 380
        for j, it in enumerate(items):
            cx = x + j * (w + gap)
            card_png(d, cx, y, w, 95, it, "", head_color=AZUL)
            if j < len(items) - 1:
                arrow_png(d, cx + w + 8, cx + w + gap - 4, y + 48, OLIVA)
        d.text((x + 2*(w+gap) + w/2, y + 150), "...e o laço recomeça", font=font(20), fill=pil(TEXTO2), anchor="mm")

    elif k == "task":
        card_png(d, 300, top, 880, 200, "Tarefa", s["items"][0], head_color=GRENA, border=GRENA)

    elif k == "summary":
        items = s["items"]; x = 300; w = 270; gap = 35; y = top
        for j, it in enumerate(items):
            cx = x + j * (w + gap)
            d.ellipse([cx + w/2 - 26, y, cx + w/2 + 26, y + 52], fill=pil(LIMA))
            d.text((cx + w/2, y + 26), "*", font=font(30, True), fill=pil(AZUL), anchor="mm")
            card_png(d, cx, y + 70, w, 120, it, "", head_color=AZUL)

    img.save(os.path.join(THUMBS, f"slide-{i:02d}.png"))

# ============================================================
# RENDER 2 — PPTX (python-pptx), editavel
# ============================================================
def box(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, wrap_=True):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h)); tf = tb.text_frame
    tf.word_wrap = wrap_; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = "Open Sans"
    r.font.color.rgb = rgb(color)
    return tb

def shape(slide, kind, x, y, w, h, fill, line=None, line_w=1.0):
    sp = slide.shapes.add_shape(kind, IN(x), IN(y), IN(w), IN(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    if line:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def card_pptx(slide, x, y, w, h, head, body, head_color=AZUL, border=None):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, BRANCO,
          line=border or "D8D8D8", line_w=2.5 if border else 0.75)
    box(slide, x + 18, y + 12, w - 36, 50, head, 18, head_color, bold=True)
    if body:
        box(slide, x + 18, y + 52, w - 36, h - 60, body, 13, TEXTO)

def brand_pptx(slide):
    shape(slide, MSO_SHAPE.RECTANGLE, 40, 92, 1200, 556, PAINEL)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 40, 40, 210, 110, OLIVA)
    shape(slide, MSO_SHAPE.RECTANGLE, 144, 150, 3, 486, OLIVA)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0, 650, 250, 70, AZUL)
    shape(slide, MSO_SHAPE.RECTANGLE, 700, 689, 320, 2, OLIVA)
    box(slide, 1030, 672, 210, 30, "cefor.ifes.edu.br", 13, OLIVA, bold=True, align=PP_ALIGN.RIGHT)

def render_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for s in SLIDES:
        sl = prs.slides.add_slide(blank); k = s["kind"]

        if k == "title":
            shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 40, 40, 210, 110, OLIVA)
            shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 0, 560, 360, 160, AZUL)
            box(sl, 120, 230, 1040, 200, s["assertion"], 40, AZUL, bold=True)
            box(sl, 120, 455, 1040, 50, s["sub"], 16, TEXTO2)
            box(sl, 700, 672, 540, 30, "cefor.ifes.edu.br", 13, OLIVA, bold=True, align=PP_ALIGN.RIGHT)
        elif k == "end":
            shape(sl, MSO_SHAPE.RECTANGLE, 0, 0, 1280, 720, AZUL)
            box(sl, 140, 300, 1000, 80, s["assertion"], 34, BRANCO, bold=True, align=PP_ALIGN.CENTER)
            box(sl, 140, 390, 1000, 50, s["sub"], 18, LIMA, align=PP_ALIGN.CENTER)
        else:
            brand_pptx(sl)
            box(sl, 290, 110, 910, 130, s["assertion"], 26, AZUL, bold=True)
            top = 290
            if k == "timeline":
                items = s["items"]; x = 300; w = 250; gap = 70; y = 360
                for j, it in enumerate(items):
                    cx = x + j*(w+gap); late = (j == len(items)-1)
                    card_pptx(sl, cx, y, w, 110, it, "", head_color=GRENA if late else AZUL,
                              border=GRENA if late else None)
                    if j < len(items)-1:
                        shape(sl, MSO_SHAPE.RIGHT_ARROW, cx+w+12, y+40, gap-24, 20, OLIVA)
            elif k == "compare":
                (h1,b1),(h2,b2) = s["items"]
                card_pptx(sl, 300, top, 410, 230, h1, b1, head_color=AZUL)
                card_pptx(sl, 770, top, 410, 230, h2, b2, head_color=GRENA)
                box(sl, 710, top+95, 60, 50, "x", 26, TEXTO2, bold=True, align=PP_ALIGN.CENTER)
            elif k == "process":
                items = s["items"]; x = 300; w = 270; gap = 35; y = top
                for j,(h1,b1) in enumerate(items):
                    cx = x + j*(w+gap)
                    shape(sl, MSO_SHAPE.OVAL, cx, y, 44, 44, OLIVA)
                    box(sl, cx, y+6, 44, 36, str(j+1), 20, BRANCO, bold=True, align=PP_ALIGN.CENTER)
                    card_pptx(sl, cx, y+60, w, 150, h1, b1, head_color=AZUL)
            elif k == "cycle":
                items = s["items"]; x = 300; w = 200; gap = 33; y = 380
                for j,it in enumerate(items):
                    cx = x + j*(w+gap)
                    card_pptx(sl, cx, y, w, 95, it, "", head_color=AZUL)
                    if j < len(items)-1:
                        shape(sl, MSO_SHAPE.RIGHT_ARROW, cx+w+6, y+33, gap-12, 20, OLIVA)
            elif k == "task":
                card_pptx(sl, 300, top, 880, 200, "Tarefa", s["items"][0], head_color=GRENA, border=GRENA)
            elif k == "summary":
                items = s["items"]; x = 300; w = 270; gap = 35; y = top
                for j,it in enumerate(items):
                    cx = x + j*(w+gap)
                    shape(sl, MSO_SHAPE.OVAL, cx+w/2-26, y, 52, 52, LIMA)
                    card_pptx(sl, cx, y+70, w, 120, it, "", head_color=AZUL)

        # notas do apresentador (Mayer: notas != tela)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]

    out = os.path.join(HERE, "avaliacao-formativa.pptx")
    prs.save(out); return out

# ============================================================
THUMBS = os.path.join(HERE, "thumbs")
os.makedirs(THUMBS, exist_ok=True)

if __name__ == "__main__":
    for i, s in enumerate(SLIDES, 1):
        render_png(i, s)
    out = render_pptx()
    print("PPTX:", out)
    print("PNGs:", THUMBS, "->", len(SLIDES), "slides")
