#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Gera o deck de exemplo "Mediacao pedagogica na EaD" no design system Cefor v2,
deck MISTO: capa/divisorias/encerramento em Degrade (Versao B) e conteudo em
Solida (Versao A). Estrutura Assertion-Evidence (titulo=afirmacao, corpo=evidencia).

Linguagem visual v2 (design-system/CEFOR-Design-System.dc.html):
  - Lima #B0CB1F como cor primaria (paineis, barras, realces).
  - Azul institucional #2C459A em titulos, seta CEFOR e numeros de secao.
  - Oliva #8C9A0D no rodape / cefor.ifes.edu.br.
  - Seta CEFOR (path oficial), logo IF/ES (grid verde + circulo vermelho),
    marca d'agua de setas, cantos arredondados e malha de degrade.

Saidas (mesma fonte de dados):
  - mediacao-pedagogica-ead.pptx   (python-pptx, editavel)  -> entregavel
  - thumbs/slide-NN.png            (Pillow)                 -> verificacao/figuras

Coordenadas em px sobre um canvas 1280x720; convertidas para polegadas no pptx.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
W, H = 1280, 720  # canvas de referencia (16:9)

# --- Paleta Cefor v2 ---
LIMA   = "B0CB1F"   # primaria
LIMA_T = "3A4400"   # texto sobre lima
AZUL   = "2C459A"
NAVY_D = "1F3A52"   # azul-escuro (titulos sobre degrade)
OLIVA  = "8C9A0D"
VERDE  = "2E9B30"
GRENA  = "CC1111"
TEXTO  = "2B2B2B"
TEXTO2 = "5A5A5A"
PAINEL = "EAEAEC"
BRANCO = "FFFFFF"
# paradas do degrade (topo-dir -> base-esq)
G_LIMA, G_VERDE, G_CIANO, G_AZUL = "C6D452", "7FC24A", "36B4A6", "3F93CE"

def rgb(h): return RGBColor.from_string(h)
def pil(h): return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

SCALE_IN = 13.333 / W
def IN(px): return Inches(px * SCALE_IN)

# --- seta CEFOR (path oficial, viewBox 0 0 100 100) ---
ARROW = [(41.8, 26.2), (73.8, 58.2), (81.6, 50.4), (82, 82),
         (50.4, 81.6), (58.2, 73.8), (26.2, 41.8)]
def arrow_pts(x, y, size):
    s = size / 100.0
    return [(x + px * s, y + py * s) for px, py in ARROW]

# --- logo IF/ES (grid, viewBox 0 0 31 42) ---
LOGO_RECTS = [(0, 11), (0, 22), (0, 33), (11, 0), (11, 11),
              (11, 22), (11, 33), (22, 0), (22, 22)]

# --- fontes Open Sans / Arial (Windows) ---
def font(size, bold=False, extra=False):
    cand = []
    if extra:
        cand += [r"C:\Windows\Fonts\OpenSans-ExtraBold.ttf"]
    if bold or extra:
        cand += [r"C:\Windows\Fonts\OpenSans-Bold.ttf", r"C:\Windows\Fonts\arialbd.ttf"]
    cand += [r"C:\Windows\Fonts\OpenSans-Regular.ttf", r"C:\Windows\Fonts\OpenSans.ttf",
             r"C:\Windows\Fonts\arial.ttf"]
    for p in cand:
        try:
            return ImageFont.truetype(p, size)
        except Exception:
            continue
    return ImageFont.load_default()

# ============================================================
# DADOS DOS SLIDES (assertion-evidence)
# ============================================================
SLIDES = [
    {"kind": "cover", "assertion": "Mediação pedagógica: estar presente faz o aluno permanecer",
     "sub": "Formação docente Cefor   ·   educação a distância",
     "notes": "Apresenta a sessão e o objetivo: sair daqui mediando, não apenas respondendo."},
    {"kind": "timeline", "assertion": "Na EaD, o silêncio do professor é lido como abandono",
     "items": ["Matrícula", "Semanas em silêncio", "Evasão"],
     "notes": "Provoca a experiência do tutor: a ausência de presença antecede a desistência."},
    {"kind": "compare", "assertion": "Transmitir conteúdo informa; mediar faz o aluno pensar junto",
     "items": [("Transmitir", "Entrega o conteúdo e espera"),
               ("Mediar", "Provoca, escuta e acompanha")],
     "notes": "Conteúdo disponível não é aprendizagem; a mediação é o que move o pensamento."},
    {"kind": "divider", "num": "01", "assertion": "O que muda quando se media",
     "notes": "Abre o bloco de demonstração das três presenças."},
    {"kind": "process", "assertion": "A presença docente se faz em três frentes ao mesmo tempo",
     "items": [("Presença de ensino", "desenha o percurso e orienta"),
               ("Presença social", "o aluno se sente entre pessoas reais"),
               ("Presença cognitiva", "constrói significado: do problema à resolução")],
     "notes": "Comunidade de Inquirição (Garrison, Anderson & Archer, 2000)."},
    {"kind": "cycle", "assertion": "Mediar é um ciclo: provocar, escutar, devolver, avançar",
     "items": ["Provocar", "Escutar", "Devolver", "Avançar"],
     "notes": "A mediação é um laço de diálogo, não uma resposta única."},
    {"kind": "divider", "num": "02", "assertion": "Mediar na prática",
     "notes": "Passa da teoria das presenças para a intervenção concreta."},
    {"kind": "compare", "assertion": "Responder encerra a conversa; perguntar mantém o aluno pensando",
     "items": [("Responder", "“Está correto.” — encerra"),
               ("Perguntar", "“O que mudaria se...?” — reabre")],
     "notes": "A pergunta devolve a investigação ao aluno; a afirmação a encerra."},
    {"kind": "task", "assertion": "Sua vez: transforme um “continue assim” em mediação",
     "items": ["Pegue um “parabéns, continue assim” e reescreva com uma pergunta que faça o aluno avançar. 5 min, em duplas."],
     "notes": "Atividade de 5 min, em duplas; comparar as reescritas."},
    {"kind": "summary", "assertion": "Mediar é transformar presença docente em aprendizagem",
     "items": ["Presente", "Dialógica", "Faz avançar"],
     "notes": "Retoma a ideia-chave: presença sentida vira permanência e aprendizagem."},
    {"kind": "end", "assertion": "Cefor — Formação docente",
     "sub": "cefor.ifes.edu.br",
     "notes": "Encerramento; logo IF/ES e canais."},
]

# ============================================================
# RENDER 1 — PNG (Pillow)
# ============================================================
def wrap(draw, text, fnt, max_w):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=fnt) <= max_w:
            cur = t
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines

def grad_img(w, h):
    """Malha diagonal: lima (topo-dir) -> verde -> ciano -> azul (base-esq)."""
    stops = [(0.0, pil(G_AZUL)), (0.32, pil(G_CIANO)), (0.66, pil(G_VERDE)), (1.0, pil(G_LIMA))]
    sw, sh = 96, 54
    base = Image.new("RGB", (sw, sh))
    px = base.load()
    for yy in range(sh):
        for xx in range(sw):
            t = (xx / (sw - 1)) * 0.5 + (1 - yy / (sh - 1)) * 0.5
            for k in range(len(stops) - 1):
                t0, c0 = stops[k]; t1, c1 = stops[k + 1]
                if t0 <= t <= t1:
                    f = (t - t0) / (t1 - t0)
                    px[xx, yy] = tuple(int(c0[i] + (c1[i] - c0[i]) * f) for i in range(3))
                    break
    return base.resize((w, h))

def arrow_png(d, x, y, size, color):
    d.polygon(arrow_pts(x, y, size), fill=pil(color))

def hline_arrow_png(d, x1, x2, y, color):
    d.line([x1, y, x2 - 10, y], fill=pil(color), width=6)
    d.polygon([(x2 - 12, y - 9), (x2, y), (x2 - 12, y + 9)], fill=pil(color))

def logo_png(d, x, y, s, mono=None):
    """Logo IF/ES. mono=None -> cores; mono='#hex' -> monocromatico."""
    circ = mono or GRENA
    rectc = mono or VERDE
    r = 4.5 * s
    cx, cy = x + 4.5 * s, y + 4.5 * s
    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=pil(circ))
    for rx, ry in LOGO_RECTS:
        d.rounded_rectangle([x + rx * s, y + ry * s, x + (rx + 9) * s, y + (ry + 9) * s],
                            radius=1.4 * s, fill=pil(rectc))

def logo_block_png(d, x, y, s, txt_color):
    logo_png(d, x, y, s, mono=None if txt_color == TEXTO else txt_color)
    tx = x + 31 * s + 13
    d.text((tx, y - 1), "INSTITUTO FEDERAL", font=font(17, extra=True), fill=pil(txt_color))
    d.text((tx, y + 21), "Espírito Santo", font=font(13, bold=True), fill=pil(txt_color))
    d.text((tx, y + 39), "Centro de Referência em Formação", font=font(12), fill=pil(txt_color))
    d.text((tx, y + 55), "e em Educação a Distância", font=font(12), fill=pil(txt_color))

def social_png(d, x, y):
    for i in range(3):
        cx = x + i * 44
        d.ellipse([cx, y, cx + 34, y + 34], fill=pil("1C1C1C"))
        d.ellipse([cx + 9, y + 9, cx + 25, y + 25], outline=pil(BRANCO), width=2)

def footer_png(d, with_line=True):
    if with_line:
        d.line([770, 690, 1090, 690], fill=pil(OLIVA), width=2)
    d.text((1240, 678), "cefor.ifes.edu.br", font=font(20, bold=True), fill=pil(OLIVA), anchor="ra")

def deco_grad_png(d, x0, x1, light="CFE6E0"):
    """Grafismos da Versao B sobre painel degrade: quadrados e circulos em contorno."""
    c = pil(light)
    d.rounded_rectangle([x0 + 26, 30, x0 + 90, 94], radius=15, outline=c, width=3)
    d.rounded_rectangle([x0 + 120, 84, x0 + 164, 128], radius=11, outline=c, width=3)
    d.rounded_rectangle([x0 + 30, 134, x0 + 80, 184], radius=12, outline=c, width=3)
    for r in (30, 50, 70):
        d.ellipse([x0 + 50 - r, 640 - r, x0 + 50 + r, 640 + r], outline=c, width=3)
    for i, cx in enumerate(range(x0 + 36, x0 + 150, 20)):
        d.ellipse([cx, 520 - (i % 2) * 7, cx + 7, 527 - (i % 2) * 7], fill=c)

def watermark_png(img, x0, x1, tint):
    """Marca d'agua de setas em mosaico sobre painel claro."""
    d = ImageDraw.Draw(img)
    for gy in range(70, 660, 70):
        for gx in range(x0 + 20, x1 - 30, 64):
            arrow_png(d, gx, gy, 30, tint)

# ---- frame de conteudo (Versao A — Solida) ----
SIDEBAR = 360
def content_frame_png(img, d):
    d.rounded_rectangle([-30, 0, SIDEBAR, 720], radius=30, fill=pil(PAINEL))
    watermark_png(img, 0, SIDEBAR, "E1E1E4")
    d = ImageDraw.Draw(img)
    arrow_png(d, 118, 250, 150, AZUL)          # seta CEFOR azul
    footer_png(d)
    return d

def title_png(d, text, left=410, top=58, right=1235, color=AZUL, size=34):
    fnt = font(size, bold=True)
    y = top
    for ln in wrap(d, text, fnt, right - left):
        d.text((left, y), ln, font=fnt, fill=pil(color)); y += int(size * 1.16)
    d.rounded_rectangle([left, y + 8, left + 96, y + 14], radius=3, fill=pil(LIMA))  # barra lima
    return y + 34

def card_png(d, x, y, w, h, head, body, head_color=AZUL, border=None, accent=LIMA):
    d.rounded_rectangle([x, y, x + w, y + h], radius=16, fill=pil(BRANCO),
                        outline=pil(border) if border else pil("DADADD"),
                        width=3 if border else 1)
    d.rounded_rectangle([x, y, x + 8, y + h], radius=4, fill=pil(accent))  # talao lateral
    d.text((x + 24, y + 16), head, font=font(23, bold=True), fill=pil(head_color))
    if body:
        fnt = font(18)
        yy = y + 56
        for ln in wrap(d, body, fnt, w - 48):
            d.text((x + 24, yy), ln, font=fnt, fill=pil(TEXTO)); yy += 24

def render_png(i, s):
    img = Image.new("RGB", (W, H), pil(BRANCO))
    d = ImageDraw.Draw(img)
    k = s["kind"]

    if k == "cover":
        img = grad_img(W, H)
        d = ImageDraw.Draw(img)
        # grafismos decorativos
        for (cx, cy, sz) in [(1150, 70, 80), (1010, 140, 42), (860, 200, 60), (770, 300, 48)]:
            d.rounded_rectangle([cx, cy, cx + sz, cy + sz], radius=int(sz * 0.22),
                                outline=(255, 255, 255), width=3)
        for r in (40, 66, 92):
            d.ellipse([1150 - r, 640 - r, 1150 + r, 640 + r], outline=(255, 255, 255), width=4)
        fnt = font(50, extra=True)
        y = 150
        d.text((90, 100), "TÍTULO DA APRESENTAÇÃO", font=font(15, bold=True), fill=(255, 255, 255))
        for ln in wrap(d, s["assertion"], fnt, 900):
            d.text((90, y), ln, font=fnt, fill=(255, 255, 255)); y += 58
        d.text((90, y + 16), s["sub"], font=font(20), fill=(255, 255, 255))
        logo_block_png(d, 90, 632, 1.5, BRANCO)
        img.save(os.path.join(THUMBS, f"slide-{i:02d}.png")); return

    if k == "divider":
        gp = grad_img(435, 720)
        img.paste(gp, (0, 0))
        d = ImageDraw.Draw(img)
        deco_grad_png(d, 0, 435, light="DFF0EA")
        d.text((470, 175), s["num"], font=font(110, extra=True), fill=pil("CFDDED"))
        fnt = font(40, bold=True); y = 320
        for ln in wrap(d, s["assertion"], fnt, 1190 - 470):
            d.text((470, y), ln, font=fnt, fill=pil(NAVY_D)); y += 46
        d.rounded_rectangle([470, y + 14, 540, y + 19], radius=3, fill=pil(VERDE))
        d.rounded_rectangle([548, y + 14, 578, y + 19], radius=3, fill=pil(G_AZUL))
        footer_png(d, with_line=False)
        img.save(os.path.join(THUMBS, f"slide-{i:02d}.png")); return

    if k == "end":
        gp = grad_img(435, 720)
        img.paste(gp, (0, 0))
        d = ImageDraw.Draw(img)
        deco_grad_png(d, 0, 435, light="DFF0EA")
        d.text((500, 170), "Obrigado!", font=font(56, extra=True), fill=pil(NAVY_D))
        d.rounded_rectangle([500, 250, 720, 256], radius=3, fill=pil(AZUL))
        logo_block_png(d, 500, 320, 1.4, TEXTO)
        social_png(d, 500, 430)
        d.text((1240, 678), s["sub"], font=font(20, bold=True), fill=pil(OLIVA), anchor="ra")
        img.save(os.path.join(THUMBS, f"slide-{i:02d}.png")); return

    # --- slides de conteudo (Solida) ---
    d = content_frame_png(img, d)
    cy = title_png(d, s["assertion"])
    top = max(cy + 6, 250)

    if k == "timeline":
        items = s["items"]; x = 430; w = 250; gap = 56; y = 360
        for j, it in enumerate(items):
            cx = x + j * (w + gap)
            late = (j == len(items) - 1)
            card_png(d, cx, y, w, 110, it, "", head_color=GRENA if late else AZUL,
                     border=GRENA if late else None, accent=GRENA if late else LIMA)
            if j < len(items) - 1:
                hline_arrow_png(d, cx + w + 8, cx + w + gap - 4, y + 55, TEXTO2)

    elif k == "compare":
        (h1, b1), (h2, b2) = s["items"]
        card_png(d, 430, top, 380, 220, h1, b1, head_color=AZUL)
        card_png(d, 850, top, 380, 220, h2, b2, head_color=VERDE, accent=VERDE)
        d.text((830, top + 110), "x", font=font(34, bold=True), fill=pil(TEXTO2), anchor="mm")

    elif k == "process":
        items = s["items"]; x = 430; w = 250; gap = 30; y = top
        for j, (h1, b1) in enumerate(items):
            cx = x + j * (w + gap)
            d.ellipse([cx, y, cx + 46, y + 46], fill=pil(AZUL))
            d.text((cx + 23, y + 23), str(j + 1), font=font(24, bold=True), fill=pil(BRANCO), anchor="mm")
            card_png(d, cx, y + 62, w, 165, h1, b1, head_color=AZUL)

    elif k == "cycle":
        items = s["items"]; x = 430; w = 178; gap = 30; y = 380
        for j, it in enumerate(items):
            cx = x + j * (w + gap)
            card_png(d, cx, y, w, 92, it, "", head_color=AZUL)
            if j < len(items) - 1:
                hline_arrow_png(d, cx + w + 6, cx + w + gap - 2, y + 46, LIMA)
        d.text((x + 3 * (w + gap) + w / 2, y + 150), "...e o laço recomeça",
               font=font(19), fill=pil(TEXTO2), anchor="mm")

    elif k == "task":
        card_png(d, 430, top, 800, 200, "Tarefa", s["items"][0], head_color=GRENA,
                 border=GRENA, accent=GRENA)

    elif k == "summary":
        items = s["items"]; x = 430; w = 250; gap = 30; y = top
        for j, it in enumerate(items):
            cx = x + j * (w + gap)
            d.ellipse([cx + w / 2 - 27, y, cx + w / 2 + 27, y + 54], fill=pil(LIMA))
            arrow_png(d, cx + w / 2 - 16, y + 11, 32, AZUL)
            card_png(d, cx, y + 74, w, 120, it, "", head_color=AZUL)

    img.save(os.path.join(THUMBS, f"slide-{i:02d}.png"))

# ============================================================
# RENDER 2 — PPTX (python-pptx, editavel)
# ============================================================
def box(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, wrap_=True):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h)); tf = tb.text_frame
    tf.word_wrap = wrap_; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = "Open Sans"
    r.font.color.rgb = rgb(color)
    return tb

def shape(slide, kind, x, y, w, h, fill, line=None, line_w=1.0, rot=0):
    sp = slide.shapes.add_shape(kind, IN(x), IN(y), IN(w), IN(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    if line:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    if rot:
        sp.rotation = rot
    sp.shadow.inherit = False
    return sp

def grad_shape(slide, x, y, w, h):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, IN(x), IN(y), IN(w), IN(h))
    sp.fill.gradient()
    stops = sp.fill.gradient_stops
    stops[0].position = 0.0; stops[0].color.rgb = rgb(G_VERDE)
    stops[1].position = 1.0; stops[1].color.rgb = rgb(G_AZUL)
    try:
        sp.fill.gradient_angle = 45.0
    except Exception:
        pass
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp

def cefor_arrow_pptx(slide, x, y, size, color):
    # aproximacao do simbolo: chevron rotacionado para baixo-direita
    shape(slide, MSO_SHAPE.CHEVRON, x, y, size, size * 0.86, color, rot=45)

def logo_pptx(slide, x, y, s, mono=None):
    circ = mono or GRENA
    rectc = mono or VERDE
    shape(slide, MSO_SHAPE.OVAL, x, y, 9 * s, 9 * s, circ)
    for rx, ry in LOGO_RECTS:
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + rx * s, y + ry * s, 9 * s, 9 * s, rectc)

def logo_block_pptx(slide, x, y, s, txt):
    logo_pptx(slide, x, y, s, mono=None if txt == TEXTO else txt)
    tx = x + 31 * s + 13
    box(slide, tx, y - 4, 420, 26, "INSTITUTO FEDERAL", 15, txt, bold=True)
    box(slide, tx, y + 18, 420, 22, "Espírito Santo", 11, txt, bold=True)
    box(slide, tx, y + 36, 420, 36, "Centro de Referência em Formação e em Educação a Distância",
        9, txt)

def footer_pptx(slide, with_line=True):
    if with_line:
        shape(slide, MSO_SHAPE.RECTANGLE, 770, 689, 320, 2, OLIVA)
    box(slide, 700, 672, 540, 30, "cefor.ifes.edu.br", 13, OLIVA, bold=True, align=PP_ALIGN.RIGHT)

def card_pptx(slide, x, y, w, h, head, body, head_color=AZUL, border=None, accent=LIMA):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, BRANCO,
          line=border or "DADADD", line_w=2.5 if border else 0.75)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 8, h, accent)
    box(slide, x + 20, y + 12, w - 40, 46, head, 17, head_color, bold=True)
    if body:
        box(slide, x + 20, y + 52, w - 40, h - 60, body, 12, TEXTO)

def content_frame_pptx(slide):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, -30, 0, SIDEBAR + 30, 720, PAINEL)
    cefor_arrow_pptx(slide, 110, 250, 150, AZUL)
    footer_pptx(slide)

def title_pptx(slide, text):
    box(slide, 410, 54, 825, 130, text, 25, AZUL, bold=True)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 410, 196, 96, 6, LIMA)

def render_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for s in SLIDES:
        sl = prs.slides.add_slide(blank); k = s["kind"]

        if k == "cover":
            grad_shape(sl, 0, 0, 1280, 720)
            for (cx, cy, sz) in [(1150, 70, 80), (1010, 140, 42), (860, 200, 60)]:
                shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, sz, sz, None, line=BRANCO, line_w=2.5)
            box(sl, 90, 96, 600, 24, "TÍTULO DA APRESENTAÇÃO", 13, BRANCO, bold=True)
            box(sl, 88, 140, 920, 220, s["assertion"], 38, BRANCO, bold=True)
            box(sl, 90, 372, 900, 40, s["sub"], 16, BRANCO)
            logo_block_pptx(sl, 90, 636, 1.5, BRANCO)

        elif k == "divider":
            grad_shape(sl, 0, 0, 435, 720)
            box(sl, 468, 150, 300, 150, s["num"], 80, "CFDDED", bold=True)
            box(sl, 468, 300, 720, 150, s["assertion"], 34, NAVY_D, bold=True)
            shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 470, 430, 70, 6, VERDE)
            shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 548, 430, 30, 6, G_AZUL)
            footer_pptx(sl, with_line=False)

        elif k == "end":
            grad_shape(sl, 0, 0, 435, 720)
            box(sl, 498, 150, 600, 90, "Obrigado!", 44, NAVY_D, bold=True)
            shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 500, 250, 220, 6, AZUL)
            logo_block_pptx(sl, 500, 322, 1.4, TEXTO)
            for i in range(3):
                shape(sl, MSO_SHAPE.OVAL, 500 + i * 44, 430, 34, 34, "1C1C1C")
            box(sl, 700, 672, 540, 30, s["sub"], 13, OLIVA, bold=True, align=PP_ALIGN.RIGHT)

        else:
            content_frame_pptx(sl)
            title_pptx(sl, s["assertion"])
            top = 250
            if k == "timeline":
                items = s["items"]; x = 430; w = 250; gap = 56; y = 360
                for j, it in enumerate(items):
                    cx = x + j * (w + gap); late = (j == len(items) - 1)
                    card_pptx(sl, cx, y, w, 110, it, "", head_color=GRENA if late else AZUL,
                              border=GRENA if late else None, accent=GRENA if late else LIMA)
                    if j < len(items) - 1:
                        shape(sl, MSO_SHAPE.RIGHT_ARROW, cx + w + 8, y + 42, gap - 12, 22, TEXTO2)
            elif k == "compare":
                (h1, b1), (h2, b2) = s["items"]
                card_pptx(sl, 430, top, 380, 220, h1, b1, head_color=AZUL)
                card_pptx(sl, 850, top, 380, 220, h2, b2, head_color=VERDE, accent=VERDE)
                box(sl, 808, top + 90, 46, 50, "x", 26, TEXTO2, bold=True, align=PP_ALIGN.CENTER)
            elif k == "process":
                items = s["items"]; x = 430; w = 250; gap = 30; y = top
                for j, (h1, b1) in enumerate(items):
                    cx = x + j * (w + gap)
                    shape(sl, MSO_SHAPE.OVAL, cx, y, 46, 46, AZUL)
                    box(sl, cx, y + 7, 46, 36, str(j + 1), 20, BRANCO, bold=True, align=PP_ALIGN.CENTER)
                    card_pptx(sl, cx, y + 62, w, 165, h1, b1, head_color=AZUL)
            elif k == "cycle":
                items = s["items"]; x = 430; w = 178; gap = 30; y = 380
                for j, it in enumerate(items):
                    cx = x + j * (w + gap)
                    card_pptx(sl, cx, y, w, 92, it, "", head_color=AZUL)
                    if j < len(items) - 1:
                        shape(sl, MSO_SHAPE.RIGHT_ARROW, cx + w + 6, y + 34, gap - 10, 22, LIMA)
            elif k == "task":
                card_pptx(sl, 430, top, 800, 200, "Tarefa", s["items"][0], head_color=GRENA,
                          border=GRENA, accent=GRENA)
            elif k == "summary":
                items = s["items"]; x = 430; w = 250; gap = 30; y = top
                for j, it in enumerate(items):
                    cx = x + j * (w + gap)
                    shape(sl, MSO_SHAPE.OVAL, cx + w / 2 - 27, y, 54, 54, LIMA)
                    card_pptx(sl, cx, y + 74, w, 120, it, "", head_color=AZUL)

        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = s["notes"]

    out = os.path.join(HERE, "mediacao-pedagogica-ead.pptx")
    prs.save(out); return out

# ============================================================
THUMBS = os.path.join(HERE, "thumbs")
os.makedirs(THUMBS, exist_ok=True)

if __name__ == "__main__":
    for i, s in enumerate(SLIDES, 1):
        render_png(i, s)
    out = render_pptx()
    print("PPTX:", out)
    print("PNGs:", THUMBS, "->", len(SLIDES), "slides")
