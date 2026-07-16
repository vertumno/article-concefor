#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract-pptx.py - Extrai todo o conteudo de um PowerPoint (.pptx).

Gera um JSON com slides (titulo, textos, TABELAS estruturadas, imagens, notas)
e salva as imagens em <saida>/assets/. Shapes AGRUPADOS sao percorridos
recursivamente; tabelas viram listas de linhas (nao sao perdidas).

Uso:
    python extract-pptx.py <entrada.pptx> [pasta_saida]

Dependencia: pip install python-pptx
"""

import json
import os
import sys


# --- Saida ASCII-safe (evita UnicodeEncodeError em console cp1252 do Windows) ---
def _log(msg):
    try:
        sys.stdout.write(msg + "\n")
    except UnicodeEncodeError:
        sys.stdout.write(msg.encode("ascii", "replace").decode("ascii") + "\n")


try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    _log("[ERRO] python-pptx nao instalado. Rode: pip install python-pptx")
    sys.exit(1)


def _walk_shapes(shapes, slide_data, assets_dir, slide_num, title_shape):
    """Percorre shapes (recursando em grupos) e acumula texto/tabelas/imagens."""
    for shape in shapes:
        # Grupo: recursao (texto/tabela/imagem dentro de grupos nao se perde)
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _walk_shapes(shape.shapes, slide_data, assets_dir, slide_num, title_shape)
            continue

        # Tabela: estrutura de linhas (antes era perdida em silencio)
        if getattr(shape, "has_table", False):
            rows = [[cell.text.strip() for cell in row.cells]
                    for row in shape.table.rows]
            slide_data["content"].append({"type": "table", "rows": rows})
            continue

        # Texto
        if shape.has_text_frame:
            if title_shape is not None and shape == title_shape:
                slide_data["title"] = shape.text
            elif shape.text.strip():
                slide_data["content"].append({"type": "text", "content": shape.text})

        # Imagem
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = shape.image
            image_name = (f"slide{slide_num}_img"
                          f"{len(slide_data['images']) + 1}.{image.ext}")
            with open(os.path.join(assets_dir, image_name), "wb") as f:
                f.write(image.blob)
            slide_data["images"].append({
                "path": f"assets/{image_name}",
                "width": shape.width,
                "height": shape.height,
            })


def extract_pptx(file_path, output_dir="."):
    """Extrai todo o conteudo do .pptx; retorna a lista de slides."""
    prs = Presentation(file_path)
    slides_data = []

    assets_dir = os.path.join(output_dir, "assets")
    os.makedirs(assets_dir, exist_ok=True)

    for slide_num, slide in enumerate(prs.slides, 1):
        slide_data = {
            "number": slide_num,
            "title": "",
            "content": [],
            "images": [],
            "notes": "",
        }

        _walk_shapes(slide.shapes, slide_data, assets_dir, slide_num,
                     slide.shapes.title)

        if slide.has_notes_slide:
            slide_data["notes"] = slide.notes_slide.notes_text_frame.text

        slides_data.append(slide_data)

    return slides_data


if __name__ == "__main__":
    if len(sys.argv) < 2:
        _log("Uso: python extract-pptx.py <entrada.pptx> [pasta_saida]")
        sys.exit(1)

    input_file = sys.argv[1]
    output_dir = sys.argv[2] if len(sys.argv) > 2 else "."
    if not os.path.isfile(input_file):
        _log(f"[ERRO] Arquivo nao encontrado: {input_file}")
        sys.exit(1)

    slides = extract_pptx(input_file, output_dir)

    output_path = os.path.join(output_dir, "extracted-slides.json")
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(slides, f, indent=2, ensure_ascii=False)

    _log(f"[OK] {len(slides)} slide(s) extraidos para {output_path}")
    for s in slides:
        tbl = sum(1 for c in s["content"] if c["type"] == "table")
        extras = f" — {len(s['images'])} imagem(ns)" + (f", {tbl} tabela(s)" if tbl else "")
        _log(f"  Slide {s['number']}: {s['title'] or '(sem titulo)'}{extras}")
