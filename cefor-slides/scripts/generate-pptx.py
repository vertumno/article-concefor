#!/usr/bin/env python3
"""
generate-pptx.py - Converter apresentacoes HTML Cefor para PowerPoint (.pptx)

Uso:
    python generate-pptx.py <entrada.html> [saida.pptx]

Dependencia:
    pip install python-pptx    # MIT license (mesma lib do extract-pptx.py)

Quando usar: o formato RECOMENDADO institucionalmente e o LibreOffice Impress
(.odp, generate-odp.py — politica de software livre do governo federal). O
.pptx existe para os casos em que e NECESSARIO (destinatario so usa PowerPoint,
modelo exigido por terceiros). python-pptx e software livre (MIT): a skill
continua sem dependencias proprietarias.

O script:
1. Le o HTML (estrutura de slides, conteudo, cores de :root).
2. Extrai blocos semanticos via deck_parser.py (titulo, antetitulo, corpo,
   bullets, KPIs .kpi-num/.kpi-label e tabelas <table> — estruturados).
3. Monta um .pptx 16:9 com a paleta Cefor (lima/azul/oliva) e Open Sans.
   Tabelas HTML viram TABELAS PowerPoint reais (cabecalho lima); KPIs viram
   linhas "numero — rotulo".
4. Valida o arquivo reabrindo-o antes de declarar sucesso; avisa se um slide
   nao contribuir nenhum conteudo (nada e perdido em silencio).

Limitacoes conscientes (o HTML continua sendo a fonte da verdade):
- Vetores decorativos (SVG, grafismos, degrades, molduras) NAO sao transpostos:
  pertencem ao layout pixel-perfect do HTML, nao ao conteudo editavel.
- Imagens <img> raster locais (png/jpg) sao embutidas; SVG e remotas sao
  puladas com aviso.
- Open Sans precisa estar instalada na maquina que abrir o .pptx (senao o
  PowerPoint substitui a fonte).
"""

import os
import re
import sys

# Parser compartilhado (mesmo do generate-odp.py) — ver deck_parser.py
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deck_parser import parse_html, log as _log, css_hex  # noqa: E402

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Cm, Pt
except ImportError:
    _log("[ERRO] python-pptx nao instalado. Rode: pip install python-pptx")
    sys.exit(1)

FONT = "Open Sans"
FOOTER_TEXT = "cefor.ifes.edu.br"

# Palco 16:9 em cm (mesma geometria do generate-odp: 33.867 x 19.05)
PAGE_W, PAGE_H = 33.867, 19.05
MARGIN_X, BODY_W = 2.2, 29.4


def _rgb(hexstr):
    return RGBColor.from_string(hexstr.lstrip("#"))


def _para(tf, first_used):
    """Primeiro paragrafo do text_frame ou um novo (evita paragrafo vazio inicial)."""
    if not first_used[0]:
        first_used[0] = True
        return tf.paragraphs[0]
    return tf.add_paragraph()


def _run(p, text, size, color, bold=False):
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    return r


def build_pptx(slides, css_root, output_path, src_dir):
    lime = _rgb(css_hex(css_root, "lime", "#B0CB1F"))
    navy = _rgb(css_hex(css_root, "navy", "#2C459A"))
    olive = _rgb(css_hex(css_root, "olive", "#8C9A0D"))
    ink = _rgb(css_hex(css_root, "ink", "#2B2B2B"))
    white = _rgb("#FFFFFF")

    prs = Presentation()
    prs.slide_width = Cm(PAGE_W)
    prs.slide_height = Cm(PAGE_H)
    blank = prs.slide_layouts[6]  # layout em branco

    def add_textbox(slide, x, y, w, h):
        box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        box.text_frame.word_wrap = True
        return box.text_frame

    def add_table(slide, tblock, y_cm, idx):
        rows = tblock["rows"]
        ncols = max(len(r) for r in rows)
        h = min(1.05 * len(rows) + 0.3, 10.5)
        gf = slide.shapes.add_table(len(rows), ncols, Cm(MARGIN_X), Cm(y_cm),
                                    Cm(BODY_W), Cm(h))
        tbl = gf.table
        tbl.first_row = False   # desliga o estilo zebrado default do PowerPoint
        tbl.horz_banding = False
        for ri, row in enumerate(rows):
            is_hdr = tblock["has_header"] and ri == 0
            for ci in range(ncols):
                cell = tbl.cell(ri, ci)
                cell.fill.solid()
                cell.fill.fore_color.rgb = lime if is_hdr else white
                p = cell.text_frame.paragraphs[0]
                _run(p, row[ci] if ci < len(row) else "",
                     14 if is_hdr else 13, ink, bold=is_hdr)

    for idx, slide_data in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        blocks = slide_data["blocks"]
        title_block = next((b for b in blocks if b[0] == "title"), None)
        kicker_block = next((b for b in blocks if b[0] == "kicker"), None)

        # Cabecalho: antetitulo + titulo
        head_tf = add_textbox(slide, MARGIN_X, 1.3, BODY_W, 4.5)
        head_used = [False]
        if kicker_block:
            _run(_para(head_tf, head_used), kicker_block[1].upper(), 14, olive, bold=True)
        if title_block:
            _run(_para(head_tf, head_used), title_block[1], 34, navy, bold=True)

        # Corpo: bullets, paragrafos e KPIs (fluxo) + tabelas (shapes proprios)
        body_tf = add_textbox(slide, MARGIN_X, 6.4, BODY_W, 10.8)
        body_used = [False]
        used = {id(title_block), id(kicker_block)}
        had_body = False
        had_flow = False
        tables = []
        kpi_num = None  # numero de KPI aguardando o rotulo seguinte
        for b in blocks:
            if id(b) in used:
                continue
            role, content = b
            if role == "table":
                tables.append(content)
                had_body = True
                continue
            if role == "kpi-num":
                kpi_num = content
                continue
            if role == "kpi-label":
                p = _para(body_tf, body_used)
                if kpi_num:
                    _run(p, kpi_num, 22, navy, bold=True)
                    _run(p, "  —  ", 18, ink)
                    kpi_num = None
                _run(p, content, 18, ink)
                had_body = had_flow = True
                continue
            if role == "bullet":
                p = _para(body_tf, body_used)
                _run(p, "▸  ", 18, lime, bold=True)
                _run(p, content, 18, ink)
                had_body = had_flow = True
            elif role in ("body", "title", "kicker"):
                _run(_para(body_tf, body_used), content, 18, ink)
                had_body = had_flow = True
        if kpi_num:  # numero sem rotulo (raro): nao perder
            _run(_para(body_tf, body_used), kpi_num, 18, ink)
            had_body = had_flow = True

        # Tabelas reais abaixo do fluxo de texto (ou no topo do corpo, se so ha tabela)
        if tables:
            y = 8.6 if had_flow else 6.4
            if len(tables) > 1:
                _log(f"[AVISO] Slide {idx}: {len(tables)} tabelas — posicione-as manualmente no PowerPoint")
            for t_i, tblock in enumerate(tables):
                add_table(slide, tblock, y + t_i * 0.8, idx)

        # Imagens raster locais (python-pptx preserva a proporcao com so a largura)
        for src in slide_data.get("images", []):
            path = os.path.join(src_dir, src)
            if not os.path.isfile(path):
                _log(f"[AVISO] Slide {idx}: imagem nao encontrada, pulada: {src}")
                continue
            try:
                pic = slide.shapes.add_picture(path, Cm(MARGIN_X), Cm(6.4), width=Cm(18.0))
                if pic.height > Cm(10.0):  # limita a altura preservando proporcao
                    ratio = Cm(10.0) / pic.height
                    pic.height = Cm(10.0)
                    pic.width = int(pic.width * ratio)
                had_body = True
            except Exception as e:  # noqa: BLE001
                _log(f"[AVISO] Slide {idx}: falha ao embutir {src}: {e}")

        # Rodape institucional
        foot_tf = add_textbox(slide, MARGIN_X, 17.8, BODY_W, 0.8)
        _run(foot_tf.paragraphs[0], FOOTER_TEXT, 11, olive)

        if not title_block and not kicker_block and not had_body:
            _log(f"[AVISO] Slide {idx}: nenhum conteudo extraido (sem classes semanticas "
                 f"nem '.reveal') — placeholder inserido; revise o slide no PowerPoint")
            _run(_para(body_tf, body_used), f"(Slide {idx})", 18, ink)

    prs.save(output_path)
    return output_path


def validate_pptx(path):
    """Reabre o .pptx para garantir que e um OOXML valido (nao apenas um ZIP)."""
    try:
        prs = Presentation(path)
        return True, len(prs.slides._sldIdLst)
    except Exception as e:  # noqa: BLE001
        return False, str(e)


# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) < 2:
        _log(__doc__)
        sys.exit(1)

    src = sys.argv[1]
    output = sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.html?$", "", src) + ".pptx"

    _log(f"[INFO] Lendo: {src}")
    slides, css_root = parse_html(src)
    if not slides:
        _log("[ERRO] Nenhum slide (<section class=\"slide\">) encontrado no HTML.")
        sys.exit(1)

    _log(f"[INFO] Slides encontrados: {len(slides)}")
    _log(f"[INFO] Cores Cefor: lima={css_root.get('lime')} azul={css_root.get('navy')} oliva={css_root.get('olive')}")

    _log("[INFO] Gerando .pptx via python-pptx (MIT)...")
    try:
        build_pptx(slides, css_root, output,
                   src_dir=os.path.dirname(os.path.abspath(src)))
    except Exception as e:  # noqa: BLE001
        _log(f"[ERRO] Falha ao gerar .pptx: {e}")
        sys.exit(1)

    ok, info = validate_pptx(output)
    if ok:
        _log(f"[OK] PPTX valido: {output} ({info} slide(s) confirmados na releitura)")
        _log("[OK] Abra no PowerPoint/Impress e revise o conteudo editavel.")
        _log("[NOTA] O formato recomendado institucionalmente e o .odp (software livre).")
    else:
        _log(f"[ERRO] Arquivo gerado NAO e um PPTX valido: {info}")
        sys.exit(1)


if __name__ == "__main__":
    main()
